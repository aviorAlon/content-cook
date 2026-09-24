"""绘本接口（契约 §6）：CRUD / 编号 / lint 问题卡 / 三个后台任务（假来源、假合成器、假装配）/ 产物流与穿越 / 上传 / 朗读稿 / 删除归档。"""

from __future__ import annotations

import base64
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from picturebook import assembly as A
from picturebook import narration as N
from webui import deps
from webui.app import app
from webui.routers import picturebook as R
from tests.picturebook_fixtures import example_data, png_bytes, wait_task
from tests.test_picturebook_pages import FakeProvider

client = TestClient(app)
MP3 = b"ID3" + b"\x00" * 64


@pytest.fixture
def project(tmp_path, monkeypatch):
    monkeypatch.setenv("STORYSTUDIO_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("STORYSTUDIO_SETTINGS", str(tmp_path / "home" / "settings.json"))
    monkeypatch.setattr(deps, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(R, "SYNTHESIZER", None)
    monkeypatch.setattr(R, "PROVIDER_FACTORY", None)
    monkeypatch.setattr(R, "DISPATCH", None)
    (tmp_path / "content" / "picturebook").mkdir(parents=True)
    return tmp_path


def _create(body=None) -> dict:
    r = client.post("/api/picturebook", json=body or example_data())
    assert r.status_code == 200, r.text
    return r.json()


def _fake_synth(profile, text, destination: Path):
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(MP3 + text.encode("utf-8"))


def _narrate(book_id: str, monkeypatch) -> dict:
    monkeypatch.setattr(R, "SYNTHESIZER", _fake_synth)
    monkeypatch.setattr(N, "probe_duration", lambda p, **kw: 1.25)
    r = client.post(f"/api/picturebook/{book_id}/narration/generate", json={})
    assert r.status_code == 200, r.text
    return wait_task(client, r.json()["task_id"])


def _illustrate(book_id: str, project: Path, monkeypatch) -> dict:
    cache = project / "cache"
    cache.mkdir(exist_ok=True)
    monkeypatch.setattr(R, "PROVIDER_FACTORY", lambda: FakeProvider(cache))
    r = client.post(f"/api/picturebook/{book_id}/pages/generate", json={})
    assert r.status_code == 200, r.text
    return wait_task(client, r.json()["task_id"])


def test_next_id_create_list_and_presets(project):
    nid = client.get("/api/picturebook/next-id").json()["id"]
    assert nid.startswith("PB-") and nid.endswith("-1")
    created = _create()
    assert created["id"] == nid and created["module"] == "picturebook" and created["warnings"] == [] and created["problems"] == []
    assert created["spec"]["pages"][0]["id"] == "P01" and "output_root" not in created["spec"]
    assert created["timeline"]["estimated"] is True and created["timeline"]["total"] > 0
    second = _create()
    assert second["id"].endswith("-2")
    assert client.post("/api/picturebook", json={**example_data(), "id": nid}).status_code == 409
    assert client.post("/api/picturebook", json={**example_data(), "id": "bad id"}).status_code == 400
    listing = client.get("/api/picturebook").json()
    assert [b["id"] for b in listing] == [nid, second["id"]]
    assert listing[0]["title"] == "小狐狸的三只袜子" and listing[0]["page_count"] == 10 and listing[0]["progress"]["pages"] == 0 and listing[0]["cover_url"] is None
    presets = client.get("/api/picturebook/presets").json()
    assert [s["name"] for s in presets["styles"]] == ["水彩", "蜡笔", "剪纸", "扁平矢量"] and presets["examples"] == ["fox_three_socks"]
    assert presets["age_bands"] == ["2-4", "4-6", "6-8"]


def test_get_put_and_validation_rejects_without_writing(project):
    book_id = _create()["id"]
    path = project / "content" / "picturebook" / f"{book_id}.json"
    before = path.read_text(encoding="utf-8")
    bad = json.loads(before)
    bad["pages"][1]["characters"] = ["ghost"]
    r = client.put(f"/api/picturebook/{book_id}", json=bad)
    assert r.status_code == 400 and "改动没通过校验，没有写盘" in r.json()["detail"] and path.read_text(encoding="utf-8") == before
    good = json.loads(before)
    good["pages"][2]["hook"] = "树洞里会不会有怪物？"
    r = client.put(f"/api/picturebook/{book_id}", json=good)
    assert r.status_code == 200
    assert any("恐怖词「怪物」" in w for w in r.json()["warnings"]) and r.json()["problems"][0]["scope"] == "page"
    assert client.put(f"/api/picturebook/{book_id}", json={**good, "id": "PB-OTHER"}).status_code == 400
    assert client.get("/api/picturebook/PB-NOPE").status_code == 404
    detail = client.get(f"/api/picturebook/{book_id}").json()
    assert detail["spec"]["pages"][2]["hook"] == "树洞里会不会有怪物？" and detail["pages"][0]["image_url"] is None
    lint = client.get(f"/api/picturebook/{book_id}/lint").json()
    assert lint["problems"][0]["target"] == "P03" and lint["problems"][0]["kind"] == "book" and lint["notes"]


def test_example_create_is_clean(project):
    r = client.post("/api/picturebook/examples/create", json={"name": "fox_three_socks"})
    assert r.status_code == 200, r.text
    assert r.json()["warnings"] == [] and r.json()["example"] == "fox_three_socks" and (project / r.json()["path"]).is_file()
    assert client.post("/api/picturebook/examples/create", json={"name": "nope"}).status_code == 404


def test_pages_task_with_fake_provider_and_artifacts(project, monkeypatch):
    book_id = _create()["id"]
    task = _illustrate(book_id, project, monkeypatch)
    assert task["status"] == "succeeded" and task["module"] == "picturebook" and task["kind"] == "pb-pages", task
    assert len(task["result"]["generated"]) == 10 and task["result"]["failed"] == []
    assert (project / "output" / "picturebook" / book_id / "05_pages" / "P01-1.png").is_file()
    detail = client.get(f"/api/picturebook/{book_id}").json()
    assert detail["pages"][0]["image_url"] == f"/api/picturebook/{book_id}/artifact/page/P01"
    assert client.get(detail["pages"][0]["image_url"]).status_code == 200
    assert client.get(f"/api/picturebook/{book_id}/artifact/page/P99").status_code == 404
    assert client.get(f"/api/picturebook/{book_id}/artifact/page/..%2F..%2Fspec").status_code == 404
    assert client.get(f"/api/picturebook/{book_id}/artifact/audio/..%2FP01_zh").status_code == 404
    assert client.get(f"/api/picturebook/{book_id}/artifact/audio/P01_zh").status_code == 404   # 还没朗读
    # 第二次：全部跳过
    again = _illustrate(book_id, project, monkeypatch)
    assert again["result"]["skipped"] and not again["result"]["generated"]
    assert client.post(f"/api/picturebook/{book_id}/pages/generate", json={"only": ["P77"]}).status_code == 404
    assert client.get("/api/picturebook").json()[0]["cover_url"] == f"/api/picturebook/{book_id}/artifact/page/P01"
    assert [t["kind"] for t in client.get("/api/tasks", params={"module": "picturebook"}).json()] == ["pb-pages", "pb-pages"]
    assert client.get("/api/tasks", params={"module": "drama"}).json() == []


def test_narration_task_with_fake_synth(project, monkeypatch):
    book_id = _create()["id"]
    task = _narrate(book_id, monkeypatch)
    assert task["status"] == "succeeded", task
    assert len(task["result"]["synthesized"]) == 20 and task["result"]["placeholder"] == {"zh": True, "en": True}
    detail = client.get(f"/api/picturebook/{book_id}").json()
    zh = detail["pages"][0]["audio"]["zh"]
    assert zh["done"] and zh["duration"] == 1.25 and zh["url"] == f"/api/picturebook/{book_id}/artifact/audio/P01_zh"
    assert client.get(zh["url"]).status_code == 200
    assert client.get(detail["pages"][3]["audio"]["zh"]["url"]).status_code == 200   # 两句页：P04_zh_1
    assert detail["timeline"]["estimated"] is False
    assert client.post(f"/api/picturebook/{book_id}/narration/generate", json={"languages": ["fr"]}).status_code == 400


def test_assemble_gates_and_task(project, monkeypatch):
    book_id = _create()["id"]
    r = client.post(f"/api/picturebook/{book_id}/assemble", json={})
    assert r.status_code == 425 and "没有插画" in r.json()["detail"]
    _illustrate(book_id, project, monkeypatch)
    assert client.post(f"/api/picturebook/{book_id}/assemble", json={}).status_code == 425
    _narrate(book_id, monkeypatch)

    def fake_assemble(spec, **kw):
        target = A.final_path(spec)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"\x00" * 2048)
        return {"final": {"path": str(target)}, "landscape": None, "total": 42.0}

    monkeypatch.setattr(A, "assemble_book", fake_assemble)
    r = client.post(f"/api/picturebook/{book_id}/assemble", json={})
    assert r.status_code == 200
    task = wait_task(client, r.json()["task_id"])
    assert task["status"] == "succeeded" and task["kind"] == "pb-assemble"
    assert client.get(f"/api/picturebook/{book_id}/artifact/final").status_code == 200
    assert client.get(f"/api/picturebook/{book_id}/artifact/final", params={"variant": "16x9"}).status_code == 404
    summary = client.get(f"/api/picturebook/{book_id}/summary").json()
    assert summary["progress"] == {"pages": 10, "pages_total": 10, "narration": 10, "narration_total": 10, "final": True, "landscape": False}
    assert summary["next"]["key"] == "done" and summary["final_url"]


def test_summary_next_steps(project, monkeypatch):
    book_id = _create()["id"]
    assert client.get(f"/api/picturebook/{book_id}/summary").json()["next"]["label"] == "生成插画"
    _illustrate(book_id, project, monkeypatch)
    assert client.get(f"/api/picturebook/{book_id}/summary").json()["next"]["label"] == "生成朗读"
    _narrate(book_id, monkeypatch)
    assert client.get(f"/api/picturebook/{book_id}/summary").json()["next"]["label"] == "装配成片"
    empty = _create({"title_zh": "空书", "pages": []})
    assert client.get(f"/api/picturebook/{empty['id']}/summary").json()["next"]["label"] == "起草分页表"


def test_script_sheet_xlsx(project):
    from openpyxl import load_workbook

    book_id = _create()["id"]
    r = client.get(f"/api/picturebook/{book_id}/script-sheet", params={"lang": "zh"})
    assert r.status_code == 200 and f"{book_id}_zh_script.xlsx" in r.headers["content-disposition"]
    path = project / "output" / "picturebook" / book_id / "09_export" / f"{book_id}_zh_script.xlsx"
    assert path.is_file()
    wb = load_workbook(path)
    assert wb.sheetnames == ["朗读稿"]
    rows = list(wb["朗读稿"].iter_rows(values_only=True))
    assert list(rows[0]) == ["页", "说话人", "文本", "语气提示"]
    assert rows[1][0] == "P01" and rows[1][1] == "旁白" and rows[1][2] == "小狐狸的三只袜子" and len(rows) == 11
    assert rows[1][3] and all(len(row) == 4 for row in rows)
    en = client.get(f"/api/picturebook/{book_id}/script-sheet", params={"lang": "en"})
    en_path = path.parent / f"{book_id}_en_script.xlsx"
    assert en.status_code == 200 and en_path.is_file()
    en_rows = list(load_workbook(en_path)["朗读稿"].iter_rows(values_only=True))
    assert list(en_rows[0]) == ["Page", "Speaker", "Text", "Delivery"]
    assert en_rows[1][:3] == ("P01", "Narrator", "Little Fox's Three Socks") and len(en_rows) == 11
    assert client.get(f"/api/picturebook/{book_id}/script-sheet", params={"lang": "fr"}).status_code == 400
    # 只有中文朗读的书没有英文稿
    zh_only = _create({**example_data(), "bilingual_mode": "zh_only"})["id"]
    refused = client.get(f"/api/picturebook/{zh_only}/script-sheet", params={"lang": "en"})
    assert refused.status_code == 400 and "没有英文朗读稿" in refused.json()["detail"]


def test_upload_assets_never_overwrite_and_update_spec(project, monkeypatch):
    # BGM / 音效的样本是只有文件头的假 mp3：按「本机没有 ffprobe、只看文件头」的退路跑（有 ffprobe 时的内容校验见 test_upload_content）
    from webui import uploads

    monkeypatch.setattr(uploads, "_tool", lambda name: None)
    book_id = _create()["id"]
    payload = {"filename": "fox.png", "data_base64": base64.b64encode(png_bytes()).decode()}
    r = client.post(f"/api/picturebook/{book_id}/assets/character/fox", json=payload)
    assert r.status_code == 200, r.text
    file = r.json()["file"]
    assert file.startswith(f"content/assets/{book_id}/characters/fox-") and (project / file).is_file()
    assert client.get(f"/api/picturebook/{book_id}").json()["spec"]["characters"]["fox"]["asset"] == file
    assert client.get(r.json()["url"]).status_code == 200
    # 同内容再传：同名不覆盖；不同内容：新文件
    assert client.post(f"/api/picturebook/{book_id}/assets/character/fox", json=payload).json()["file"] == file
    other = {"filename": "fox.png", "data_base64": base64.b64encode(png_bytes((1, 2, 3))).decode()}
    assert client.post(f"/api/picturebook/{book_id}/assets/character/fox", json=other).json()["file"] != file
    assert len(list((project / "content" / "assets" / book_id / "characters").glob("*.png"))) == 2
    music = {"filename": "lullaby.mp3", "data_base64": base64.b64encode(MP3).decode()}
    r = client.post(f"/api/picturebook/{book_id}/assets/music/bgm", json=music)
    assert r.status_code == 200 and client.get(f"/api/picturebook/{book_id}").json()["spec"]["music"]["file"] == r.json()["file"]
    r = client.post(f"/api/picturebook/{book_id}/assets/sfx/flip", json=music)
    assert client.get(f"/api/picturebook/{book_id}").json()["spec"]["page_turn_sfx"] == r.json()["file"]
    assert client.post(f"/api/picturebook/{book_id}/assets/character/ghost", json=payload).status_code == 404
    assert client.post(f"/api/picturebook/{book_id}/assets/poster/x", json=payload).status_code == 400
    assert client.post(f"/api/picturebook/{book_id}/assets/character/fox", json={"filename": "fox.exe", "data_base64": "AA=="}).status_code == 400
    assert client.get(f"/api/picturebook/{book_id}/artifact/asset/character/..%2Fspec.json").status_code == 404


def test_delete_archives_spec_and_output(project, monkeypatch):
    book_id = _create()["id"]
    _illustrate(book_id, project, monkeypatch)
    r = client.delete(f"/api/picturebook/{book_id}")
    assert r.status_code == 200 and r.json()["output_moved"] is True
    archived = project / r.json()["archived_to"]
    assert (archived / "spec.json").is_file() and (archived / "output" / "05_pages" / "P01-1.png").is_file()
    assert not (project / "content" / "picturebook" / f"{book_id}.json").exists()
    assert not (project / "output" / "picturebook" / book_id).exists()
    assert client.delete(f"/api/picturebook/{book_id}").status_code == 404 and client.get("/api/picturebook").json() == []
