"""朗读稿（契约 §6）：``09_export/<书 id>_<zh|en>_script.xlsx``，交给真人朗读或别的 TTS 用。

一页一行，工作表「朗读稿」。中文稿的表头是 ``页 / 说话人 / 文本 / 语气提示``，英文稿是 ``Page / Speaker / Text / Delivery``。
「页」就是页 id：录好的音频按 ``P01_zh.mp3`` / ``P01_en.mp3`` 放回 ``07_audio``，``narration.sync_imported`` 会认领成整页一段。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

from .spec import export_dir

SHEET_NAME = "朗读稿"
HEADERS = {"zh": ("页", "说话人", "文本", "语气提示"), "en": ("Page", "Speaker", "Text", "Delivery")}
DEFAULT_SPEAKER = {"zh": "旁白", "en": "Narrator"}
DELIVERY_BY_AGE = {
    "zh": {"2-4": "温柔、慢速、睡前故事腔，字字清楚", "4-6": "温柔、亲切的讲故事语气，节奏从容", "6-8": "自然的讲故事语气，情绪随文本起伏"},
    "en": {"2-4": "gentle, slow, bedtime-story voice, every word clear", "4-6": "warm storytelling voice, unhurried",
           "6-8": "natural storytelling voice, expressive but calm"},
}
_COLUMN_WIDTHS = (10, 16, 60, 40)


class ScriptSheetError(ValueError):
    pass


def sheet_filename(spec: Mapping[str, Any], lang: str) -> str:
    return f"{spec.get('id')}_{lang}_script.xlsx"


def sheet_rows(spec: Mapping[str, Any], lang: str) -> list[dict[str, Any]]:
    """每页一行：``{page, speaker, text, delivery}``；这一种语言没有文字的页跳过。"""

    if lang not in HEADERS:
        raise ScriptSheetError("lang 只能是 zh / en")
    mode = str(spec.get("bilingual_mode") or "")
    if (lang == "zh" and mode == "en_only") or (lang == "en" and mode == "zh_only"):
        raise ScriptSheetError(f"这本书的双语策略是 {mode}，没有{'中文' if lang == 'zh' else '英文'}朗读稿")
    narration = spec.get("narration") or {}
    speaker = str(narration.get(f"{lang}_voice") or "").strip() or DEFAULT_SPEAKER[lang]
    delivery = DELIVERY_BY_AGE[lang].get(str(spec.get("age_band") or "2-4"), DELIVERY_BY_AGE[lang]["2-4"])
    rows: list[dict[str, Any]] = []
    for page in spec.get("pages") or ():
        text = str(page.get(lang) or "").strip()
        if text:
            rows.append({"page": str(page["id"]), "speaker": speaker, "text": text, "delivery": delivery})
    return rows


def write_script_sheet(spec: Mapping[str, Any], lang: str, destination: Path | None = None) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ScriptSheetError("openpyxl 没装：pip install openpyxl") from exc

    rows = sheet_rows(spec, lang)
    target = destination or (export_dir(spec) / sheet_filename(spec, lang))
    target.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append(list(HEADERS[lang]))
    for row in rows:
        sheet.append([row["page"], row["speaker"], row["text"], row["delivery"]])
    for index, width in enumerate(_COLUMN_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    book.save(str(target))
    return target


__all__ = ["SHEET_NAME", "HEADERS", "DEFAULT_SPEAKER", "ScriptSheetError", "sheet_filename", "sheet_rows", "write_script_sheet"]
